"""Full-text exclusion decided by the coders.

Each coder records, for a study, whether the full text must be excluded and why.
Decisions live one file per coder in ``Screening/fulltext - <coder>.json``.  When
``AGREE_MIN`` distinct coders agree on exclusion the study is marked in the master
screening workbook (``Fulltext_Verdict = Exclude`` + ``Fulltext_Exclusion_Criterion``)
and listed in ``Screening/fulltext_agreed.json``; the coding tab hides it and the
analysis pipeline drops it.

The canonical agreed list is also what ``Scripts/study_exclusions.py`` reads so the
exclusion reaches ``metaanalysis.py`` / ``reanalysis.py`` / ``publication_bias.py``.
"""
from datetime import datetime
import glob
import json
import os
import re
import shutil
import tempfile
import threading

from studio_lib import atomic_replace, safe_token
from studio_state import S, log

# Controlled exclusion criteria (mirrors the values used in master_dataset.xlsx).
FULLTEXT_CRITERIA = [
    {"value": "C1", "label": "No misinformation intervention (C1)"},
    {"value": "C2", "label": "No control/comparison group (C2)"},
    {"value": "C3", "label": "Not English (C3)"},
    {"value": "C4", "label": "Not a primary experimental study (C4)"},
    {"value": "irretrievable", "label": "Fulltext irretrievable"},
    {"value": "no_contrast", "label": "No eligible contrast reported"},
    {"value": "duplicate", "label": "Duplicate/non-independent report of an included study"},
    {"value": "observational", "label": "No experimental intervention (observational)"},
    {"value": "domain_mismatch", "label": "Outcome/intervention outside misinformation susceptibility"},
    {"value": "other", "label": "Other (specify)"},
]
_CRIT_LABEL = {c["value"]: c["label"] for c in FULLTEXT_CRITERIA}

AGREE_MIN = 2

_MARK_LOCK = threading.Lock()
_MARK_STATE = {"running": False, "last": None, "marked": [], "error": "", "finished": None}

_MAP_CACHE = {"mtime": None, "map": None, "info": None}


def criteria():
    return {"criteria": FULLTEXT_CRITERIA, "agree_min": AGREE_MIN}


# ---------------------------------------------------------------- per-coder store
def decision_path(coder):
    return os.path.join(S.SCREENING, f"fulltext - {safe_token(coder)}.json")


def load_decisions(coder):
    if not coder or not str(coder).strip():
        return {}
    path = decision_path(coder)
    if not os.path.exists(path):
        return {}
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def save_decisions(coder, decisions):
    os.makedirs(S.SCREENING, exist_ok=True)
    path = decision_path(coder)
    fd, tmp = tempfile.mkstemp(suffix=".json", dir=S.TEMP)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(decisions, fh, ensure_ascii=False, indent=1, sort_keys=True)
            fh.write("\n")
        atomic_replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


def all_decisions():
    """{study_key: {coder: decision}} across every coder file."""
    out = {}
    for path in glob.glob(os.path.join(S.SCREENING, "fulltext - *.json")):
        base = os.path.basename(path)[: -len(".json")]
        coder = base[len("fulltext - "):]
        for key, dec in load_decisions(coder).items():
            if isinstance(dec, dict):
                out.setdefault(str(key), {})[coder] = dec
    return out


def save_decision(coder, key, exclude, criterion, note):
    """Store one coder's decision and refresh the agreed list."""
    coder = str(coder or "").strip()
    key = str(key or "").strip()
    if not coder:
        raise ValueError("coder name is required")
    if not key:
        raise ValueError("study key is required")
    decisions = load_decisions(coder)
    if exclude:
        crit = str(criterion or "other").strip() or "other"
        decisions[key] = {
            "exclude": True,
            "criterion": crit,
            "criterion_label": _CRIT_LABEL.get(crit, crit),
            "note": str(note or "").strip(),
            "saved_at": datetime.now().isoformat(timespec="seconds"),
        }
    else:
        decisions[key] = {
            "exclude": False, "criterion": "", "criterion_label": "",
            "note": str(note or "").strip(),
            "saved_at": datetime.now().isoformat(timespec="seconds"),
        }
    save_decisions(coder, decisions)
    return recompute_agreed()


def compute_agreed():
    """Studies excluded by >= AGREE_MIN distinct coders (majority criterion)."""
    per = all_decisions()
    agreed = {}
    for key, by_coder in per.items():
        excl = {c: d for c, d in by_coder.items() if d.get("exclude")}
        if len(excl) < AGREE_MIN:
            continue
        counts = {}
        for d in excl.values():
            counts[d.get("criterion") or "other"] = counts.get(d.get("criterion") or "other", 0) + 1
        crit = max(counts, key=counts.get)
        notes = [d.get("note") for d in excl.values() if d.get("note")]
        agreed[key] = {
            "criterion": crit,
            "criterion_label": _CRIT_LABEL.get(crit, crit),
            "coders": sorted(excl.keys()),
            "notes": notes,
            "n_exclude": len(excl),
        }
    return agreed


def agreed_path():
    return os.path.join(S.SCREENING, "fulltext_agreed.json")


def load_agreed():
    path = agreed_path()
    if not os.path.exists(path):
        return {}
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def save_agreed(agreed):
    os.makedirs(S.SCREENING, exist_ok=True)
    path = agreed_path()
    fd, tmp = tempfile.mkstemp(suffix=".json", dir=S.TEMP)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(agreed, fh, ensure_ascii=False, indent=1, sort_keys=True)
            fh.write("\n")
        atomic_replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


def recompute_agreed(mark=True):
    """Recompute the agreed list and mark newly-agreed studies in the master."""
    agreed = compute_agreed()
    previous = load_agreed()
    now = datetime.now().isoformat(timespec="seconds")
    merged = {}
    for key, info in agreed.items():
        prior = previous.get(key) or {}
        merged[key] = dict(info)
        merged[key]["marked_at"] = prior.get("marked_at") or now
        merged[key]["master"] = prior.get("master") or ""
    save_agreed(merged)
    newly = [k for k in merged if not (previous.get(k) or {}).get("master")]
    if mark and newly:
        start_marking(newly)
    return merged


def start_marking(keys):
    with _MARK_LOCK:
        if _MARK_STATE["running"]:
            return False
        _MARK_STATE.update(running=True, error="", finished=None, marked=[])
    threading.Thread(target=_mark_worker, args=(list(keys),), daemon=True).start()
    return True


def mark_status():
    return {k: v for k, v in _MARK_STATE.items()}


# ---------------------------------------------------------------- master workbook
def _norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def _norm_doi(s):
    return re.sub(r"^https?://(dx\.)?doi\.org/", "", str(s or "").strip().lower())


def _extraction_info():
    info = {}
    for path in glob.glob(os.path.join(S.EXTRACTION, "*.json")):
        key = os.path.splitext(os.path.basename(path))[0]
        try:
            with open(path, encoding="utf-8") as fh:
                data = json.load(fh)
        except Exception:
            continue
        rec = data[0] if isinstance(data, list) and data else (data if isinstance(data, dict) else {})
        if isinstance(rec, dict):
            info[key] = (_norm_doi(rec.get("DOI")), _norm(rec.get("Title")))
    return info


def master_mapping():
    """{study_key: (row_index)} into master_dataset.xlsx (DOI, then filename, then title)."""
    import openpyxl

    master = os.path.join(S.ROOT, "master_dataset.xlsx")
    if not os.path.exists(master):
        return {}
    mtime = int(os.path.getmtime(master))
    if _MAP_CACHE["map"] is not None and _MAP_CACHE["mtime"] == mtime:
        return _MAP_CACHE["map"]
    info = _extraction_info()
    by_doi, by_title, by_fn = {}, {}, {}
    wb = openpyxl.load_workbook(master, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it)]
        H = {h: i for i, h in enumerate(head)}

        def cell(row, name):
            i = H.get(name)
            if i is None or i >= len(row) or row[i] is None:
                return ""
            return row[i]

        for idx, row in enumerate(it, start=2):
            doi = _norm_doi(cell(row, "doi"))
            title = _norm(cell(row, "title"))
            fn = str(cell(row, "PDF_Filename") or "").strip()
            stem = os.path.splitext(os.path.basename(fn))[0] if fn else ""
            if doi:
                by_doi.setdefault(doi, idx)
            if title:
                by_title.setdefault(title, idx)
            if stem:
                by_fn.setdefault(stem, idx)
    finally:
        wb.close()
    mapping = {}
    for key, (doi, title) in info.items():
        row = None
        if doi and doi in by_doi:
            row = by_doi[doi]
        elif key in by_fn:
            row = by_fn[key]
        elif title and title in by_title:
            row = by_title[title]
        if row is not None:
            mapping[key] = row
    _MAP_CACHE.update(mtime=mtime, map=mapping, info=info)
    return mapping


def _mark_worker(keys):
    state = _MARK_STATE
    try:
        _mark_master(keys)
    except Exception as exc:  # noqa: BLE001
        state.update(error=str(exc))
        log(f"fulltext mark failed: {exc}")
    finally:
        state.update(running=False, finished=datetime.now().isoformat(timespec="seconds"))


def _mark_master(keys):
    import openpyxl

    master = os.path.join(S.ROOT, "master_dataset.xlsx")
    mapping = master_mapping()
    agreed = load_agreed()
    marked, missing = [], []
    if not os.path.exists(master):
        _MARK_STATE.update(error="master_dataset.xlsx not found")
        return
    backup_dir = os.path.join(S.ROOT, "archive", "coding_app_ui_backups",
                              "fulltext_merge_" + datetime.now().strftime("%Y%m%d_%H%M%S"))
    os.makedirs(backup_dir, exist_ok=True)
    shutil.copy2(master, os.path.join(backup_dir, os.path.basename(master)))

    wb = openpyxl.load_workbook(master)
    try:
        ws = wb.worksheets[0]
        head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        H = {h: i + 1 for i, h in enumerate(head)}   # 1-based column indexes

        def setcell(row, name, value):
            col = H.get(name)
            if col is not None:
                ws.cell(row=row, column=col, value=value)

        for key in keys:
            info = agreed.get(key) or {}
            row = mapping.get(key)
            if not row:
                missing.append(key)
                continue
            setcell(row, "Fulltext_Verdict", "Exclude")
            setcell(row, "Fulltext_Exclusion_Criterion", info.get("criterion_label") or "Coder agreement")
            note = "; ".join(info.get("notes") or []) or f"Excluded by coders: {', '.join(info.get('coders') or [])}"
            setcell(row, "Fulltext_Reason", note)
            if "Effective_Verdict" in H:
                setcell(row, "Effective_Verdict", "Exclude")
            marked.append(key)
        tmp = os.path.join(S.TEMP, "master_dataset.fulltext.xlsx")
        wb.save(tmp)
        wb.close()
        atomic_replace(tmp, master)
    except Exception:
        wb.close()
        raise
    # record marking outcome on the agreed entries
    cur = load_agreed()
    now = datetime.now().isoformat(timespec="seconds")
    for key in keys:
        if key in cur:
            cur[key]["master"] = "marked" if key in marked else "not_found"
            cur[key]["marked_at"] = now
    save_agreed(cur)
    _MAP_CACHE.update(mtime=None, map=None)
    _MARK_STATE.update(marked=marked)
    log(f"fulltext merge marked={len(marked)} missing={len(missing)}")
