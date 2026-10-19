"""APA 7th-edition (PRISMA-based) report assembly for the Coding Studio
analysis export.

The Flask app collects the data (``ctx``) with its own helpers and calls
``build_docx`` / ``build_markdown`` / ``prisma_png`` here.  This module has no
Flask or project imports, so it is safe to bundle in the packaged app; it only
needs ``python-docx`` (optional) and ``Pillow`` (for the PRISMA chart).
"""
import os


# ---------------------------------------------------------------- PRISMA chart
def _font(size):
    from PIL import ImageFont
    for name in ("DejaVuSans.ttf", "arial.ttf", "segoeui.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


def prisma_png(counts, path):
    """Draw the PRISMA 2020 flow diagram (two-column layout, black and white).

    Left track: database/register searches.  Right track: records identified via
    other sources go straight to full-text screening.  Full-text exclusion
    reasons are the named review criteria (criteria with n = 0 are omitted).
    """
    import math
    from PIL import Image, ImageDraw
    W = 1490
    img = Image.new("RGB", (W, 2400), "white")
    d = ImageDraw.Draw(img)
    f_title = _font(24)
    f_head = _font(19)
    f_box = _font(20)      # box headlines
    f_small = _font(18)    # other text (one size smaller)
    f_phase = _font(18)

    def wrap(text, font, maxw):
        out = []
        for para in str(text).split("\n"):
            words = para.split()
            if not words:
                out.append("")
                continue
            cur = ""
            for w in words:
                t = (cur + " " + w).strip()
                if d.textlength(t, font=font) <= maxw:
                    cur = t
                else:
                    if cur:
                        out.append(cur)
                    cur = w
            if cur:
                out.append(cur)
        return out

    def box(x, y, w, h, text, font=f_box, thick=False):
        d.rectangle([x, y, x + w, y + h], outline="black", width=3 if thick else 2, fill="white")
        lines = wrap(text, font, w - 18)
        lh = font.size + 6
        cy = y + (h - len(lines) * lh) / 2
        for ln in lines:
            tw = d.textlength(ln, font=font)
            d.text((x + (w - tw) / 2, cy), ln, font=font, fill="black")
            cy += lh

    def panel_lines(w, items):
        return [wrap(f"{label} (n = {n:,})", f_small, w - 30)
                for label, n in items if n]

    def panel_h(ilines):
        return 18 + (f_box.size + 6) + sum(len(g) for g in ilines) * (f_small.size + 6)

    def draw_panel(x, y, w, title, ilines):
        h = panel_h(ilines)
        d.rectangle([x, y, x + w, y + h], outline="black", width=2, fill="white")
        d.text((x + 14, y + 9), title, font=f_box, fill="black")
        cy = y + 9 + f_box.size + 6
        for g in ilines:
            for ln in g:
                d.text((x + 24, cy), ln, font=f_small, fill="black")
                cy += f_small.size + 6
        return h

    def arrow(x1, y1, x2, y2):
        d.line([x1, y1, x2, y2], fill="black", width=2)
        ang = math.atan2(y2 - y1, x2 - x1)
        L, Wd = 14, 6
        p2 = (x2 - L * math.cos(ang) + Wd * math.sin(ang),
              y2 - L * math.sin(ang) - Wd * math.cos(ang))
        p3 = (x2 - L * math.cos(ang) - Wd * math.sin(ang),
              y2 - L * math.sin(ang) + Wd * math.cos(ang))
        d.polygon([(x2, y2), p2, p3], fill="black")

    def vtext(cx, cy, text, font):
        tmp = Image.new("RGBA", (1, 1))
        w = int(ImageDraw.Draw(tmp).textlength(text, font=font)) + 8
        h = font.size + 8
        timg = Image.new("RGBA", (w, h), (255, 255, 255, 0))
        ImageDraw.Draw(timg).text((4, 4), text, font=font, fill="black")
        timg = timg.rotate(90, expand=True)
        img.paste(timg, (int(cx - timg.width / 2), int(cy - timg.height / 2)), timg)

    def group_items(T):
        gg = T.get("excluded_groups") or {}
        return [(g, gg.get(g, 0)) for g in (counts.get("ft_groups") or [])]

    t = counts.get("tracks") or {}
    db = t.get("database") or {}
    oth = t.get("other") or {}
    coded = counts.get("included", 0)

    d.text((28, 14), "PRISMA 2020 flow diagram", font=f_title, fill="black")

    # columns
    lx, bw = 40, 250
    px, pw = 300, 430
    ox, ow = 742, 360
    opx, opw = 1114, 370
    bh = 96
    cx = lx + bw // 2
    ocx = ox + ow // 2

    # ---- column headers ----
    hy, hh = 54, 48
    d.rectangle([lx, hy, px + pw, hy + hh], outline="black", width=2, fill="white")
    d.text((lx + 12, hy + 13), "Identification of studies via databases and registers",
           font=f_head, fill="black")
    d.rectangle([ox, hy, opx + opw, hy + hh], outline="black", width=2, fill="white")
    d.text((ox + 12, hy + 13), "Identification of studies via other methods",
           font=f_head, fill="black")

    # ---- panel text (measured before layout) ----
    ar_i = panel_lines(pw, [(k, v) for k, v in (db.get("abstract_reasons") or {}).items()])
    ft_i = panel_lines(pw, group_items(db))
    oft_i = panel_lines(opw, group_items(oth))

    y_id = 118
    y_scr = 286
    y_assess = y_scr + max(bh, panel_h(ar_i)) + 36
    y_inc = y_assess + max(bh, panel_h(ft_i), panel_h(oft_i)) + 36

    # ---- left track ----
    box(lx, y_id, bw, bh, f"Records identified\n(n = {db.get('identified', 0):,})")
    box(lx, y_scr, bw, bh, f"Records screened\n(n = {db.get('screened', 0):,})")
    box(lx, y_assess, bw, bh, f"Reports assessed for eligibility\n(n = {db.get('assessed', 0):,})")
    arrow(cx, y_id + bh, cx, y_scr)
    arrow(cx, y_scr + bh, cx, y_assess)
    draw_panel(px, y_scr, pw, "Records excluded", ar_i)
    arrow(lx + bw, y_scr + bh // 2, px, y_scr + bh // 2)
    draw_panel(px, y_assess, pw, "Records excluded", ft_i)
    arrow(lx + bw, y_assess + bh // 2, px, y_assess + bh // 2)

    # ---- right track (sources straight to full text) ----
    mapping = {"Meta-analyses / reviews": "reviews", "Other sources": "author recommendations"}
    src_lines = []
    for k, v in (oth.get("sources") or {}).items():
        src_lines += wrap(f"{mapping.get(k, k.lower())} (n = {v:,})", f_small, ow - 24)
    src_h = 16 + (f_box.size + 6) + len(src_lines) * (f_small.size + 6)
    d.rectangle([ox, y_id, ox + ow, y_id + src_h], outline="black", width=2, fill="white")
    cy = y_id + 10
    head = "Records identified from"
    tw = d.textlength(head, font=f_box)
    d.text((ox + (ow - tw) / 2, cy), head, font=f_box, fill="black")
    cy += f_box.size + 6
    for ln in src_lines:
        tw = d.textlength(ln, font=f_small)
        d.text((ox + (ow - tw) / 2, cy), ln, font=f_small, fill="black")
        cy += f_small.size + 6
    arrow(ocx, y_id + src_h, ocx, y_assess)
    box(ox, y_assess, ow, bh, f"Reports assessed for eligibility\n(n = {oth.get('assessed', 0):,})")
    draw_panel(opx, y_assess, opw, "Records excluded", oft_i)
    arrow(ox + ow, y_assess + bh // 2, opx, y_assess + bh // 2)

    # ---- included (receives both tracks) ----
    arrow(cx, y_assess + bh, cx, y_inc)
    box(lx, y_inc, bw, bh, f"Studies included\n(n = {coded:,})", thick=True)
    d.line([ocx, y_assess + bh, ocx, y_inc + bh // 2], fill="black", width=2)
    arrow(ocx, y_inc + bh // 2, lx + bw, y_inc + bh // 2)

    # ---- phase labels ----
    vtext(28, (hy + y_id + bh) // 2, "Identification", f_phase)
    vtext(28, (y_scr + y_assess + bh) // 2, "Screening", f_phase)
    vtext(28, y_inc + bh // 2, "Included", f_phase)

    used = y_inc + bh + 30
    img.crop((0, 0, W, min(2400, used))).save(path, "PNG")
    return path


# ---------------------------------------------------------------- docx helpers
def _docx():
    try:
        import docx
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        return docx, Pt, Inches, WD_ALIGN_PARAGRAPH, qn, OxmlElement
    except Exception:
        return None


def _set_cell_borders(cell, top=None, bottom=None):
    pass


def _style_table(docx_mod, table):
    """APA-style table: no vertical rules, rule under the header and at the end."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_border(el, edge, sz=6, color="000000"):
        tcPr = el.get_or_add_tcPr()
        borders = tcPr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcPr.append(borders)
        tag = qn("w:" + edge)
        b = borders.find(tag)
        if b is None:
            b = OxmlElement("w:" + edge)
            borders.append(b)
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), color)

    for r_i, row in enumerate(table.rows):
        for cell in row.cells:
            if r_i == 0:
                set_border(cell._tc, "bottom", 8)
            if r_i == len(table.rows) - 1:
                set_border(cell._tc, "bottom", 8)


def build_docx(ctx, path):
    d = _docx()
    if d is None:
        return None
    docx, Pt, Inches, WD_ALIGN_PARAGRAPH, qn, OxmlElement = d
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(0)
    try:
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    except Exception:
        pass
    try:
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
    except Exception:
        pass

    def keep_next(p):
        """Word 'keep with next': do not break before the following paragraph/table."""
        try:
            p.paragraph_format.keep_with_next = True
        except Exception:
            pass

    def no_split(table):
        """Prevent a table row from splitting across a page boundary."""
        for row in table.rows:
            try:
                trPr = row._tr.get_or_add_trPr()
                if trPr.find(qn("w:cantSplit")) is None:
                    trPr.append(OxmlElement("w:cantSplit"))
            except Exception:
                pass

    def keep_table_with_note(table):
        """Keep the table on the same page as the paragraph that follows it."""
        for cell in table.rows[-1].cells:
            for pp in cell.paragraphs:
                keep_next(pp)

    def h(level, text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        if level >= 3:
            r.italic = True
        p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if level == 1
                       else WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
        keep_next(p)
        return p

    def para(text, indent=True):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        return p

    def table_block(idx, title, columns, rows, note=""):
        p = doc.add_paragraph(); r = p.add_run(f"Table {idx}"); r.bold = True
        p2 = doc.add_paragraph(); r2 = p2.add_run(title); r2.italic = True
        keep_next(p); keep_next(p2)
        cols = [str(c) for c in columns]
        t = doc.add_table(rows=1, cols=len(cols))
        for j, c in enumerate(cols):
            t.rows[0].cells[j].text = c
            for run in t.rows[0].cells[j].paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for j, c in enumerate(cols):
                v = row.get(c, "")
                cells[j].text = "" if v is None else str(v)
        for row in t.rows:
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.line_spacing = 1.0
                    pp.paragraph_format.space_after = Pt(0)
                    pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in pp.runs:
                        run.font.size = Pt(10)
        _style_table(docx, t)
        no_split(t)
        if note:
            keep_table_with_note(t)
            p3 = doc.add_paragraph(); r3 = p3.add_run("Note. "); r3.italic = True
            p3.add_run(note)
            p3.paragraph_format.line_spacing = 1.0
            p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        return t

    # ---------------- Title + Method ----------------
    title = doc.add_paragraph(); tr = title.add_run(ctx.get("title", "Methods and Results"))
    tr.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h(1, "Method")
    for sec in ctx.get("methods", []):
        h(2, sec["heading"])
        for p in sec.get("paragraphs", []):
            para(p)
        for tb in sec.get("tables", []):
            table_block(ctx["_table_no"], tb["title"], tb["columns"], tb["rows"],
                        tb.get("note", ""))
            ctx["_table_no"] += 1
        for fig in sec.get("figures", []):
            p = doc.add_paragraph(); r = p.add_run(f"Figure {ctx['_figure_no']}"); r.bold = True
            p2 = doc.add_paragraph(); r2 = p2.add_run(fig.get("title", "")); r2.italic = True
            keep_next(p); keep_next(p2)
            if fig.get("path") and os.path.exists(fig["path"]):
                doc.add_picture(fig["path"], width=Inches(6.4))
            ctx["_figure_no"] += 1

    # ---------------- Results ----------------
    h(1, "Results")
    for sec in ctx.get("results", []):
        h(2, sec["heading"])
        for p in sec.get("paragraphs", []):
            para(p)
        for tb in sec.get("tables", []):
            table_block(ctx["_table_no"], tb["title"], tb["columns"], tb["rows"],
                        tb.get("note", ""))
            ctx["_table_no"] += 1
        for fig in sec.get("figures", []):
            p = doc.add_paragraph(); r = p.add_run(f"Figure {ctx['_figure_no']}"); r.bold = True
            p2 = doc.add_paragraph(); r2 = p2.add_run(fig.get("title", "")); r2.italic = True
            keep_next(p); keep_next(p2)
            if fig.get("path") and os.path.exists(fig["path"]):
                doc.add_picture(fig["path"], width=Inches(6.4))
            ctx["_figure_no"] += 1

    doc.save(path)
    return path


# ---------------------------------------------------------------- markdown
def build_markdown(ctx):
    lines = [f"# {ctx.get('title', 'Methods and Results')}", ""]
    lines.append(f"*Generated {ctx.get('generated', '')}.*")
    lines.append("")
    for level, text in ctx.get("outline", []):
        lines.append("#" * min(level, 6) + " " + text)
    for sec in ctx.get("methods", []):
        lines.append(f"## {sec['heading']}")
        for p in sec.get("paragraphs", []):
            lines.append(p)
        for tb in sec.get("tables", []):
            lines.append(f"**{tb['title']}**")
            cols = [str(c) for c in tb["columns"]]
            lines.append("| " + " | ".join(cols) + " |")
            lines.append("| " + " | ".join("---" for _ in cols) + " |")
            for row in tb["rows"]:
                lines.append("| " + " | ".join(
                    "" if row.get(c) is None else str(row.get(c)) for c in cols) + " |")
            lines.append("")
        for fig in sec.get("figures", []):
            lines.append(f"**{fig.get('title','Figure')}** — `{fig.get('path','')}`")
            lines.append("")
        lines.append("")
    lines.append("# Results")
    for sec in ctx.get("results", []):
        lines.append(f"## {sec['heading']}")
        for p in sec.get("paragraphs", []):
            lines.append(p)
        for tb in sec.get("tables", []):
            lines.append(f"**{tb['title']}**")
            cols = [str(c) for c in tb["columns"]]
            lines.append("| " + " | ".join(cols) + " |")
            lines.append("| " + " | ".join("---" for _ in cols) + " |")
            for row in tb["rows"]:
                lines.append("| " + " | ".join(
                    "" if row.get(c) is None else str(row.get(c)) for c in cols) + " |")
            lines.append("")
        for fig in sec.get("figures", []):
            lines.append(f"**{fig.get('title','Figure')}** — `{fig.get('path','')}`")
            lines.append("")
    return "\n".join(lines)
