"""Harmonizer routes: the wide-table variable editor plus its save/missing/index endpoints."""

from datetime import datetime
from flask import Blueprint, abort, jsonify, request
import csv
import glob
import json
import os
import re
import tempfile
import threading
from studio_lib import atomic_replace, domain_vocab, safe_token
from studio_state import S

bp = Blueprint("harmonizer", __name__)


# ---------------------------------------------------------------- harmonizer tab
# Lets a coder audit every harmonized variable: its uniform label, role, domain,
# veracity, timepoint and polarity.  Variables that share a base name and differ
# only by a numeric index (e.g. ``..._i01`` .. ``..._i06``) are bundled so the
# shared parameters can be set once, while the expanded list allows per-variable
# overrides.  Overrides are written to ``Codings/<Key> - harmonizer - <coder>.json``
# and never touch the harmonized dataset or its sidecar.

_HARM_IDX = re.compile(r"(_i?\d+)+$")

_HARM_FIELDS = ("role", "dv_domain", "veracity", "timepoint", "dv_polarity")

def harmonized_paths(key):
    folder = os.path.join(S.DATASETS or "", key)
    return (os.path.join(folder, key + " - harmonized_variables.csv"),
            os.path.join(folder, key + " - harmonized.csv"))

def harmonizer_bundle(name):
    base = _HARM_IDX.sub("", str(name))
    return base or str(name)

def harmonizer_coder_files(key):
    out = []
    for f in glob.glob(os.path.join(S.CODINGS, glob.escape(key) + " - harmonizer - *.json")):
        base = os.path.basename(f)[: -len(".json")]
        coder = base[len(key) + len(" - harmonizer - "):]
        try:
            stamp = datetime.fromtimestamp(os.path.getmtime(f)).isoformat(timespec="seconds")
        except OSError:
            stamp = None
        out.append({"coder": coder, "file": os.path.basename(f), "modified": stamp})
    return sorted(out, key=lambda d: d["coder"].lower())

def load_harmonizer_overrides(key):
    """Merged overrides from every coder file (later files win), plus a per-coder map."""
    merged, by_coder = {}, {}
    for info in harmonizer_coder_files(key):
        try:
            with open(os.path.join(S.CODINGS, info["file"]), encoding="utf-8") as fh:
                data = json.load(fh)
        except Exception:
            continue
        ov = data.get("overrides") or {}
        by_coder[info["coder"]] = ov
        for uv, fields in ov.items():
            if isinstance(fields, dict):
                merged.setdefault(uv, {}).update(fields)
    return merged, by_coder

def load_harmonized_variables(key):
    vfile, wide = harmonized_paths(key)
    if not os.path.isfile(vfile):
        return None
    rows = []
    with open(vfile, encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            uv = (row.get("uniform_variable") or "").strip()
            if not uv:
                continue
            m = _HARM_IDX.search(uv)
            rows.append({
                "uniform_variable": uv,
                "role": (row.get("role") or "").strip(),
                "dv_domain": (row.get("dv_domain") or "").strip(),
                "veracity": (row.get("veracity") or "").strip(),
                "timepoint": (row.get("timepoint") or "").strip(),
                "dv_polarity": (row.get("dv_polarity") or "").strip(),
                "original_variable": (row.get("original_variable") or "").strip(),
                "artifact": (row.get("artifact") or "").strip(),
                "item_scope": (row.get("item_scope") or "").strip(),
                "notes": (row.get("notes") or "").strip(),
                "bundle": harmonizer_bundle(uv),
                "index": (m.group(0).lstrip("_") if m else ""),
            })
    return {"variables": rows, "has_wide": os.path.isfile(wide)}

def _harmonizer_known(key):
    return key in S.KEYS or os.path.isfile(harmonized_paths(key)[0])

def _harmonizer_context(key):
    """Study context for the harmonizer: the coded experiment metadata and its DV
    rows (plus the coding notes), so a coder can align the harmonized variables
    with the extraction.  Read-only; never written back."""
    from studio_lib import load_extraction
    out = []
    for rec in (load_extraction(key) or []):
        if not isinstance(rec, dict):
            continue
        dvs = []
        for dv in (rec.get("DV_Metrics") or []):
            if not isinstance(dv, dict):
                continue

            def _join(v):
                if isinstance(v, (list, tuple)):
                    return "; ".join(str(x) for x in v if x not in (None, ""))
                return str(v or "")

            dvs.append({
                "dv_name": str(dv.get("dv_name") or "")[:140],
                "dv_domain": _join(dv.get("dv_domain")),
                "score_types": _join(dv.get("score_types")),
                "timepoint": str(dv.get("timepoint_1") or ""),
                "polarity": str(dv.get("dv_polarity") or ""),
                "eligible": dv.get("eligible_primary"),
                "veracity": str(dv.get("veracity") or ""),
            })
        out.append({
            "study_id": str(rec.get("Study_ID") or ""),
            "title": str(rec.get("Title") or "")[:300],
            "description": str(rec.get("Study_Description") or "")[:800],
            "authors": str(rec.get("Authors") or ""),
            "year": rec.get("Year") or "",
            "n_dv": len(dvs),
            "dvs": dvs[:80],
        })
    notes = ""
    notes_path = os.path.join(S.EXTRACTION, key + ".md")
    if os.path.isfile(notes_path):
        try:
            with open(notes_path, encoding="utf-8", errors="replace") as fh:
                notes = fh.read()[:4000]
        except OSError:
            notes = ""
    return {"records": out, "notes": notes}

@bp.get("/api/harmonizer/<path:key>")
def api_harmonizer(key):
    if not _harmonizer_known(key):
        abort(404)
    data = load_harmonized_variables(key)
    overrides, by_coder = load_harmonizer_overrides(key)
    context = _harmonizer_context(key)
    if data is None:
        return jsonify({"key": key, "exists": False, "variables": [], "bundles": [],
                        "coders": harmonizer_coder_files(key), "overrides": {},
                        "by_coder": {}, "context": context, "options": _harmonizer_options()})
    bundles = {}
    for v in data["variables"]:
        b = bundles.setdefault(v["bundle"], {"bundle": v["bundle"], "role": v["role"],
                                             "variables": []})
        b["variables"].append(v["uniform_variable"])
    bundle_list = sorted(bundles.values(), key=lambda b: (b["role"], b["bundle"].lower()))
    return jsonify({"key": key, "exists": True, "has_wide": data["has_wide"],
                    "variables": data["variables"], "bundles": bundle_list,
                    "coders": harmonizer_coder_files(key), "overrides": overrides,
                    "by_coder": by_coder, "context": context, "options": _harmonizer_options()})

_HARM_MISSING_LOCK = threading.Lock()

_HARM_MISSING_CACHE = {"mtime": None, "data": None}

# Fields a DV must carry.  ``veracity`` and ``timepoint`` are deliberately NOT
# required: veracity only exists for true/false stimuli (it is decoded from the
# DV name and is blank otherwise), and a blank timepoint is the study's primary
# posttest.  Requiring them flagged ~7.5k of 9.4k DVs as "missing".
_HARM_REQUIRED_DV_FIELDS = ("dv_domain", "dv_polarity")

def _harmonizer_missing_map():
    """{key: {"missing": n, "dv": n_dv, "dv_missing": n_dv_missing}}.

    A variable is "missing" when a required harmonized field is empty: ``role``
    for every variable, and ``dv_domain`` + ``dv_polarity`` for a DV.  Covariates
    and other roles legitimately omit the DV fields; veracity and timepoint are
    optional for a DV (see ``_HARM_REQUIRED_DV_FIELDS``)."""
    ds = S.DATASETS or ""
    files = (glob.glob(os.path.join(ds, "*", "* - harmonized_variables.csv"))
             if os.path.isdir(ds) else [])
    newest = 0.0
    for f in files:
        try:
            newest = max(newest, os.path.getmtime(f))
        except OSError:
            pass
    with _HARM_MISSING_LOCK:
        if _HARM_MISSING_CACHE["data"] is not None and _HARM_MISSING_CACHE["mtime"] == newest:
            return _HARM_MISSING_CACHE["data"]
    out = {}
    for f in files:
        key = os.path.basename(os.path.dirname(f))
        n_missing = n_dv = n_dv_missing = 0
        try:
            with open(f, encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    if not (row.get("uniform_variable") or "").strip():
                        continue
                    role = (row.get("role") or "").strip()
                    dv_ok = all((row.get(k) or "").strip() for k in _HARM_REQUIRED_DV_FIELDS)
                    if not role:
                        n_missing += 1
                    elif role == "dv" and not dv_ok:
                        n_missing += 1
                    if role == "dv":
                        n_dv += 1
                        if not dv_ok:
                            n_dv_missing += 1
        except Exception:
            continue
        out[key] = {"missing": n_missing, "dv": n_dv, "dv_missing": n_dv_missing}
    with _HARM_MISSING_LOCK:
        _HARM_MISSING_CACHE.update(mtime=newest, data=out)
    return out

@bp.get("/api/harmonizer/missing")
def api_harmonizer_missing():
    return jsonify({"stats": _harmonizer_missing_map()})

@bp.get("/api/harmonizer/index")
def api_harmonizer_index():
    """Dataset keys that are not part of the paper list (experiment-level artefacts)."""
    extra = []
    for f in glob.glob(os.path.join(S.DATASETS or "", "*", "* - harmonized_variables.csv")):
        k = os.path.basename(os.path.dirname(f))
        if k in S.KEYS:
            continue
        extra.append({"key": k, "n_dv": 0, "n_records": 0, "n_coders": 0,
                      "missing_metrics": 0, "needs_reason": 0, "module": None,
                      "has_harmonized": True, "has_extraction": False,
                      "has_harmonizer": bool(glob.glob(os.path.join(
                          S.CODINGS, glob.escape(k) + " - harmonizer - *.json")))})
    return jsonify({"papers": extra})

def _harmonizer_options():
    return {"role": ["dv", "covariate", ""],
            "dv_domain": list(domain_vocab()) + [""],
            "veracity": ["true", "false", ""],
            "timepoint": ["pre", "post", "followup", ""],
            "dv_polarity": ["higher_is_better", "higher_is_worse", "ambiguous", ""]}

@bp.post("/api/harmonizer/save")
def api_harmonizer_save():
    payload = request.get_json(force=True, silent=True) or {}
    key = payload.get("key")
    coder = payload.get("coder")
    overrides = payload.get("overrides")
    if not _harmonizer_known(key):
        return jsonify({"ok": False, "error": "unknown paper key"}), 400
    if not coder or not str(coder).strip():
        return jsonify({"ok": False, "error": "coder name is required"}), 400
    if not isinstance(overrides, dict):
        return jsonify({"ok": False, "error": "no overrides to save"}), 400
    clean = {}
    for uv, fields in overrides.items():
        if not isinstance(fields, dict):
            continue
        keep = {k: v for k, v in fields.items() if k in _HARM_FIELDS and v is not None}
        if keep:
            clean[str(uv)] = keep
    target = os.path.join(S.CODINGS, f"{key} - harmonizer - {safe_token(coder)}.json")
    existed = os.path.exists(target)
    record = {"key": key, "coder": coder,
              "saved": datetime.now().isoformat(timespec="seconds"), "overrides": clean}
    fd, tmp = tempfile.mkstemp(suffix=".json", dir=S.TEMP)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(record, fh, ensure_ascii=False, indent=2)
            fh.write("\n")
        atomic_replace(tmp, target)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return jsonify({"ok": True, "path": os.path.relpath(target, S.ROOT).replace("\\", "/"),
                    "overwrote": existed, "overrides": len(clean), "records": len(clean)})


# ===========================================================================
# Assignment workspace (plan §4 B3): raw-variable context, the live decoder
# preview and the coder-owned per-study assignment drafts + apply.
# ===========================================================================
from contextlib import contextmanager  # noqa: E402
import io  # noqa: E402
import sys  # noqa: E402

_DRAFT_FIELDS = ("artifact", "original_variable", "source_file", "item_scope",
                 "uniform_variable", "role", "dv_domain", "veracity", "timepoint",
                 "dv_polarity", "slot_fix", "evidence", "note", "status")


def _s(v):
    """Clean string: None/NaN sentinels and blanks -> "" (bool -> true/false)."""
    if isinstance(v, bool):
        return "true" if v else "false"
    if type(v).__name__ == "bool_":
        return "true" if bool(v) else "false"
    if v is None:
        return ""
    try:
        import pandas as pd
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def _scripts_dir():
    return os.path.join(S.ROOT, "Scripts")


def _ensure_scripts_path():
    d = _scripts_dir()
    if d and d not in sys.path:
        sys.path.insert(0, d)
    return d


@contextmanager
def _quiet():
    """Silence a module import's stdout (harmonizer_template prints on import)."""
    buf = io.StringIO()
    old = sys.stdout
    sys.stdout = buf
    try:
        yield
    finally:
        sys.stdout = old


def _import_assignments():
    _ensure_scripts_path()
    import variable_assignments as va  # noqa: PLC0415
    return va


def _import_template():
    _ensure_scripts_path()
    with _quiet():
        import harmonizer_template as ht  # noqa: PLC0415
    return ht


def _import_contract():
    _ensure_scripts_path()
    import contract as ct  # noqa: PLC0415
    return ct


def _import_dvp():
    _ensure_scripts_path()
    import dv_polarity as dvp  # noqa: PLC0415
    return dvp


def _study_folder(key):
    return os.path.join(S.DATASETS or "", key)


def _known_study(key):
    return os.path.isdir(_study_folder(key))


# ------------------------------------------------------------------ decode
@bp.get("/api/harmonizer/decode")
def api_harmonizer_decode():
    """Live decode preview using the harmonizer's own decoders.

    ``domain_of``/``timepoint_of`` come from the per-study harmonizer template
    (``harmonizer_template.py``); veracity is decoded by BOTH the canonical
    ``contract.dv_veracity`` and the template's ``veracity_of`` so a name/field
    disagreement is visible, never silent.  Expected polarity uses
    ``dv_polarity.expected_polarity``.
    """
    name = (request.args.get("name") or "").strip()
    score_types = (request.args.get("score_types") or "").strip()
    notes = (request.args.get("notes") or "").strip()
    original = (request.args.get("original") or "").strip()
    out = {"name": name, "domain": "", "veracity": "", "veracity_contract": "",
           "veracity_template": "", "veracity_agreement": None, "timepoint": "",
           "polarity": "", "polarity_source": "", "errors": []}
    if not name:
        return jsonify(out)
    try:
        ht = _import_template()
        out["domain"] = ht.domain_of(name)
        out["veracity_template"] = ht.veracity_of(name)
        out["timepoint"] = ht.timepoint_of(name)
    except Exception as e:  # noqa: BLE001
        out["errors"].append(f"template decoders unavailable: {e!r}")
    try:
        ct = _import_contract()
        out["veracity_contract"] = ct.dv_veracity(
            name, score_types.split(",") if score_types else None, notes)
    except Exception as e:  # noqa: BLE001
        out["errors"].append(f"contract.dv_veracity unavailable: {e!r}")
    vt, vc = out["veracity_template"], out["veracity_contract"]
    if vt or vc:
        out["veracity_agreement"] = (vt == vc)
        out["veracity"] = vc or vt
    try:
        dvp = _import_dvp()
        pol = dvp.expected_polarity(out["domain"] or "", out["veracity"] or name,
                                    name, notes, original)
        out["polarity"] = pol or ""
        out["polarity_source"] = "dv_polarity.expected_polarity"
    except Exception as e:  # noqa: BLE001
        out["errors"].append(f"dv_polarity unavailable: {e!r}")
    return jsonify(out)


# ------------------------------------------------------------------ rawvar
def _header_of_csv(path):
    try:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return [str(c).strip().lstrip("\ufeff") for c in next(csv.reader(fh))]
    except (OSError, StopIteration):
        return []


def _series_from_csv(path, var, want_pid=False):
    import pandas as pd
    hdr = _header_of_csv(path)
    if var not in hdr:
        return None, None, None
    usecols = [var] + ([c for c in ("pid", "condition", "condition_raw")
                        if want_pid and c in hdr] if want_pid else [])
    df = pd.read_csv(path, usecols=usecols, encoding="utf-8-sig", low_memory=False)
    return df[var], (df if want_pid else None), None


def _series_from_spss(path, var):
    import pyreadstat
    reader = pyreadstat.read_dta if path.lower().endswith(".dta") else pyreadstat.read_sav
    df, meta = reader(path, usecols=[var])
    if var not in df.columns:
        return None, None, None, None
    labels = {}
    try:
        labels = dict((meta.variable_value_labels or {}).get(var) or {})
    except Exception:  # noqa: BLE001
        labels = {}
    vlabel = ""
    try:
        vlabel = (meta.column_names_to_labels or {}).get(var) or ""
    except Exception:  # noqa: BLE001
        vlabel = ""
    return df[var], labels, vlabel, meta


def _series_from_rds(path, var):
    import pyreadr
    obj = pyreadr.read_r(path)
    for _name, df in (obj or {}).items():
        if hasattr(df, "columns") and var in df.columns:
            return df[var], None, None
    return None, None, None


def _series_from_xlsx(path, var):
    try:
        import openpyxl
    except Exception:  # noqa: BLE001
        return None, None, None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            hdr = None
            for row in rows:
                if row and any(c not in (None, "") for c in row):
                    hdr = [str(c).strip() if c is not None else "" for c in row]
                    break
            if not hdr or var not in hdr:
                continue
            idx = hdr.index(var)
            vals = []
            for row in rows:
                if row and idx < len(row):
                    vals.append(row[idx])
            import pandas as pd
            return pd.Series(vals, name=var), None, ws.title
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return None, None, None


def _labels_from_spss(path, names):
    """Metadata-only value labels / variable label for the first matching name."""
    try:
        import pyreadstat
    except Exception:  # noqa: BLE001
        return {}, ""
    reader = pyreadstat.read_dta if path.lower().endswith(".dta") else pyreadstat.read_sav
    try:
        _df, meta = reader(path, metadataonly=True)
    except Exception:  # noqa: BLE001
        return {}, ""
    for nm in names:
        try:
            labels = dict((meta.variable_value_labels or {}).get(nm) or {})
        except Exception:  # noqa: BLE001
            labels = {}
        vlabel = ""
        try:
            vlabel = (meta.column_names_to_labels or {}).get(nm) or ""
        except Exception:  # noqa: BLE001
            vlabel = ""
        if labels or vlabel:
            return labels, vlabel
    return {}, ""


def _stats_and_examples(series):
    import pandas as pd
    s = pd.Series(series)
    n = int(len(s))
    nn = s.dropna()
    numeric = pd.to_numeric(nn, errors="coerce")
    num = numeric.dropna()
    stats = {
        "n": n,
        "missing": int(s.isna().sum()),
        "n_unique": int(nn.nunique()) if len(nn) else 0,
        "min": (float(num.min()) if len(num) else None),
        "max": (float(num.max()) if len(num) else None),
        "mean": (round(float(num.mean()), 4) if len(num) else None),
        "numeric": bool(len(num)),
    }
    top = []
    try:
        vc = nn.astype(str).value_counts().head(12)
        top = [{"value": str(k), "n": int(v)} for k, v in vc.items()]
    except Exception:  # noqa: BLE001
        top = []
    return stats, top


def _preview_columns(header, var, uniform, limit=12):
    """Up to ``limit`` informative columns: the target var/uniform column plus
    pid/condition when present, then the first columns in file order.  Never
    dumps hundreds of columns of a wide harmonized table."""
    present = set(header)
    lower = {str(c).lower(): c for c in header}
    important = []
    for nm in (var, uniform):
        if nm and nm in present and nm not in important:
            important.append(nm)
    for nm in ("pid", "condition", "condition_raw"):
        c = lower.get(nm)
        if c and c not in important:
            important.append(c)
    sel = list(important)
    for c in header:
        if len(sel) >= limit:
            break
        if c not in sel:
            sel.append(c)
    sel = sel[:limit]
    order = {c: i for i, c in enumerate(header)}
    sel.sort(key=lambda c: order.get(c, 1 << 30))
    return sel


def _preview_highlight(header, var, uniform):
    for nm in (var, uniform):
        if nm and nm in header:
            return nm
    return ""


def _preview_rows(df, cols, n_limit):
    """Python-native string rows (JSON-safe; numpy scalars never leak)."""
    out = []
    for row in df[cols].itertuples(index=False, name=None):
        out.append(["" if _isna(v) else _s(v) for v in row])
        if len(out) >= n_limit:
            break
    return out


def _preview_csv(path, kind, var, uniform, n_limit=10):
    import pandas as pd
    header = _header_of_csv(path)
    if not header:
        return None
    cols = _preview_columns(header, var, uniform)
    if not cols:
        return None
    n_rows = _csv_data_rows(path)
    df = pd.read_csv(path, usecols=cols, encoding="utf-8-sig", low_memory=False,
                     nrows=n_limit)
    return {"file": os.path.basename(path), "kind": kind,
            "n_rows": n_rows, "n_cols": len(header), "columns": cols,
            "rows": _preview_rows(df, cols, n_limit),
            "highlight": _preview_highlight(header, var, uniform)}


def _csv_data_rows(path):
    """Data-row count via streaming (header excluded); never loads the frame."""
    n = -1
    try:
        with open(path, encoding="utf-8-sig", newline="", errors="replace") as fh:
            for _row in csv.reader(fh):
                n += 1
    except OSError:
        return 0
    return max(n, 0)


def _preview_spss(path, kind, var, uniform, n_limit=10):
    import pyreadstat
    reader = pyreadstat.read_dta if path.lower().endswith(".dta") else pyreadstat.read_sav
    df, meta = reader(path, row_limit=n_limit)
    header = [str(c) for c in df.columns]
    cols = _preview_columns(header, var, uniform)
    if not cols:
        return None
    try:
        n_rows = int(getattr(meta, "number_rows", 0) or 0)
    except (TypeError, ValueError):
        n_rows = 0
    if n_rows <= 0:
        n_rows = int(df.shape[0])
    try:
        n_cols = int(getattr(meta, "number_columns", 0) or 0)
    except (TypeError, ValueError):
        n_cols = 0
    if n_cols <= 0:
        n_cols = len(header)
    return {"file": os.path.basename(path), "kind": kind,
            "n_rows": n_rows, "n_cols": n_cols, "columns": cols,
            "rows": _preview_rows(df, cols, n_limit),
            "highlight": _preview_highlight(header, var, uniform)}


def _preview_xlsx(path, kind, var, uniform, n_limit=10):
    try:
        import openpyxl
    except Exception:  # noqa: BLE001
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            header = None
            for row in it:
                if row and any(c not in (None, "") for c in row):
                    header = ["" if c is None else str(c).strip() for c in row]
                    break
            if not header:
                continue
            cols = _preview_columns(header, var, uniform)
            if not cols:
                continue
            idxs = [header.index(c) for c in cols]
            rows = []
            for row in it:
                if row is None:
                    break
                rows.append(["" if i >= len(row) or row[i] is None else _s(row[i])
                             for i in idxs])
                if len(rows) >= n_limit:
                    break
            try:
                n_rows = int(ws.max_row or len(rows))
            except (TypeError, ValueError):
                n_rows = len(rows)
            return {"file": os.path.basename(path), "kind": kind,
                    "n_rows": n_rows, "n_cols": len(header), "columns": cols,
                    "rows": rows, "highlight": _preview_highlight(header, var, uniform)}
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return None


def _dataset_preview(key, folder, var, uniform, artifact):
    """A small read-only preview of the study's dataset (first ~10 rows, up to
    ~12 informative columns).  Prefers the requested ``prepared`` artifact, then
    any prepared CSV, then the harmonized wide table, then the raw source.
    Returns ``None`` when nothing is readable — the UI degrades gracefully."""
    cands = []
    csvs = sorted(glob.glob(os.path.join(folder, f"{key} - prepared*.csv")))
    if artifact:
        preferred = os.path.join(folder, f"{key} - prepared {artifact}.csv")
        if os.path.isfile(preferred):
            cands.append((preferred, "prepared"))
    for c in csvs:
        if all(c != p for p, _ in cands):
            cands.append((c, "prepared"))
    harm = os.path.join(folder, f"{key} - harmonized.csv")
    if os.path.isfile(harm):
        cands.append((harm, "harmonized"))
    srcs = _rawvar_sources(folder, key)
    for ext_key, kind in (("csv", "csv"), ("sav", "sav"), ("dta", "dta"),
                          ("xlsx", "xlsx")):
        for base in srcs.get(ext_key, []):
            p = os.path.join(folder, base)
            if all(p != q for q, _ in cands):
                cands.append((p, kind))
    for path, kind in cands:
        low = path.lower()
        try:
            if low.endswith(".csv"):
                got = _preview_csv(path, kind, var, uniform)
            elif low.endswith((".sav", ".dta")):
                got = _preview_spss(path, kind, var, uniform)
            elif low.endswith((".xlsx", ".xls")):
                got = _preview_xlsx(path, kind, var, uniform)
            else:
                got = None
        except Exception:  # noqa: BLE001  (never block an assignment)
            got = None
        if got and got.get("columns"):
            return got
    return None


def _rawvar_sources(folder, key):
    out = {"csv": [], "sav": [], "dta": [], "rds": [], "xlsx": []}
    for f in sorted(glob.glob(os.path.join(folder, "*"))):
        base = os.path.basename(f)
        low = base.lower()
        if not os.path.isfile(f):
            continue
        if low.endswith(".csv"):
            out["csv"].append(base)
        elif low.endswith(".sav"):
            out["sav"].append(base)
        elif low.endswith(".dta"):
            out["dta"].append(base)
        elif low.endswith((".rds", ".rdata")):
            out["rds"].append(base)
        elif low.endswith((".xlsx", ".xls")):
            out["xlsx"].append(base)
    return out


def _snippets(key, names):
    """Small paper/coding-notes snippets mentioning the variable name."""
    out = []
    targets = [n for n in names if n and len(str(n)) >= 3]
    if not targets:
        return out
    sources = [(os.path.join(S.TEXT, key + ".txt"), "paper text"),
               (os.path.join(S.EXTRACTION, key + ".md"), "coding notes")]
    for path, label in sources:
        if not os.path.isfile(path):
            continue
        try:
            with open(path, encoding="utf-8", errors="replace") as fh:
                text = fh.read()
        except OSError:
            continue
        low = text.lower()
        for name in targets:
            i = low.find(str(name).lower())
            if i < 0:
                continue
            a, b = max(0, i - 160), min(len(text), i + 220)
            out.append({"source": label, "text": text[a:b].replace("\n", " ").strip()})
            if len(out) >= 6:
                return out
    return out


@bp.get("/api/rawvar/<path:key>")
def api_rawvar(key):
    """Read-only context for one raw variable (labels, stats, examples, snippets).

    Best-effort and **never blocks an assignment**: a missing/unreadable source
    yields an empty payload plus a ``degraded`` reason, not an error.
    """
    if not _known_study(key):
        abort(404)
    var = (request.args.get("var") or "").strip()
    artifact = (request.args.get("artifact") or "").strip()
    folder = _study_folder(key)
    out = {"key": key, "var": var, "artifact": artifact, "found": False,
           "source": None, "value_labels": [], "variable_label": "",
           "stats": None, "top_values": [], "examples": [], "example_cols": [],
           "original_values": "", "mapped_values": "", "snippets": [],
           "dataset_preview": None,
           "codebook_files": [], "sources": _rawvar_sources(folder, key),
           "degraded": []}

    # variables.csv context (original_values / mapped_values / notes)
    var_path = os.path.join(folder, f"{key} - variables.csv")
    if os.path.isfile(var_path):
        try:
            import pandas as pd
            v = pd.read_csv(var_path, encoding="utf-8-sig", low_memory=False)
            v.columns = [str(c).strip().lstrip("\ufeff") for c in v.columns]
            ucol = "uniform_variable" if "uniform_variable" in v.columns else (
                "uniform_name" if "uniform_name" in v.columns else None)
            hit = v[v["original_variable"].astype(str) == var]
            if len(hit):
                r = hit.iloc[0]
                out["original_values"] = str(r.get("original_values") or "")[:2000]
                out["mapped_values"] = str(r.get("mapped_values") or "")[:2000]
                if ucol:
                    out["uniform_variable"] = str(r.get(ucol) or "")
        except Exception as e:  # noqa: BLE001
            out["degraded"].append(f"variables.csv unreadable: {e!r}")
    else:
        out["degraded"].append("no variables.csv for this study")

    uniform = (request.args.get("uniform") or "").strip()

    # Dataset preview (item 2): a small read-only view of the prepared/harmonized
    # table so the context pane shows the DATASET, not just one column.  Built
    # before the "no variable" early return so study-scope targets get it too;
    # absent/unreadable sources degrade to None and never block an assignment.
    try:
        out["dataset_preview"] = _dataset_preview(
            key, folder, var, uniform or out.get("uniform_variable", ""), artifact)
    except Exception as e:  # noqa: BLE001
        out["dataset_preview"] = None
        out["degraded"].append(f"dataset preview failed: {e!r}")

    if not var:
        out["degraded"].append("no variable requested")
        return jsonify(out)

    names = [n for n in (var, uniform) if n]
    series = None
    # 1) prepared CSVs (prefer the requested artifact)
    csvs = sorted(glob.glob(os.path.join(folder, f"{key} - prepared*.csv")))
    ordered = []
    if artifact:
        for c in csvs:
            if os.path.basename(c) in (f"{key} - prepared {artifact}.csv",):
                ordered.insert(0, c)
    ordered += [c for c in csvs if c not in ordered]
    for c in ordered:
        for nm in names:
            try:
                s, extra, _ = _series_from_csv(c, nm, want_pid=True)
            except Exception as e:  # noqa: BLE001
                out["degraded"].append(f"{os.path.basename(c)} unreadable: {e!r}")
                break
            if s is not None:
                series = s
                out["found"] = True
                out["matched_name"] = nm
                out["source"] = {"file": os.path.basename(c), "kind": "csv"}
                if extra is not None:
                    ex = extra.head(10)
                    out["example_cols"] = [str(x) for x in ex.columns]
                    out["examples"] = [
                        {str(k): (None if _isna(val) else val) for k, val in row.items()}
                        for _, row in ex.iterrows()]
                break
        if series is not None:
            break

    # 2) SPSS / Stata
    if series is None:
        for c in (glob.glob(os.path.join(folder, "*.sav"))
                  + glob.glob(os.path.join(folder, "*.dta"))):
            for nm in names:
                try:
                    s, labels, vlabel, _meta = _series_from_spss(c, nm)
                except Exception as e:  # noqa: BLE001
                    out["degraded"].append(f"{os.path.basename(c)} unreadable: {e!r}")
                    break
                if s is not None:
                    series = s
                    out["found"] = True
                    out["matched_name"] = nm
                    out["source"] = {"file": os.path.basename(c),
                                     "kind": "sav" if c.lower().endswith(".sav") else "dta"}
                    out["variable_label"] = vlabel or ""
                    out["value_labels"] = [{"value": str(k), "label": str(v)}
                                           for k, v in (labels or {}).items()]
                    break
            if series is not None:
                break

    # 3) RDS
    if series is None:
        for c in (glob.glob(os.path.join(folder, "*.rds"))
                  + glob.glob(os.path.join(folder, "*.RData"))
                  + glob.glob(os.path.join(folder, "*.rdata"))):
            for nm in names:
                try:
                    s, _l, _v = _series_from_rds(c, nm)
                except Exception as e:  # noqa: BLE001
                    out["degraded"].append(f"{os.path.basename(c)} unreadable: {e!r}")
                    break
                if s is not None:
                    series = s
                    out["found"] = True
                    out["matched_name"] = nm
                    out["source"] = {"file": os.path.basename(c), "kind": "rds"}
                    break
            if series is not None:
                break

    # 4) XLSX
    if series is None:
        for c in (glob.glob(os.path.join(folder, "*.xlsx"))
                  + glob.glob(os.path.join(folder, "*.xls"))):
            for nm in names:
                try:
                    s, _l, sheet = _series_from_xlsx(c, nm)
                except Exception as e:  # noqa: BLE001
                    out["degraded"].append(f"{os.path.basename(c)} unreadable: {e!r}")
                    break
                if s is not None:
                    series = s
                    out["found"] = True
                    out["matched_name"] = nm
                    out["source"] = {"file": os.path.basename(c), "kind": "xlsx",
                                     "sheet": sheet}
                    break
            if series is not None:
                break

    if series is None:
        out["degraded"].append("raw column not found in any source file")
    else:
        # value labels often live only in the SPSS/Stata file while the series
        # was read from the (already renamed) prepared CSV -- fetch them from
        # metadata cheaply when still missing.
        if not out["value_labels"]:
            for c in (glob.glob(os.path.join(folder, "*.sav"))
                      + glob.glob(os.path.join(folder, "*.dta"))):
                labels, vlabel = _labels_from_spss(c, names)
                if labels or vlabel:
                    out["value_labels"] = [{"value": str(k), "label": str(v)}
                                           for k, v in labels.items()]
                    out["variable_label"] = vlabel or out["variable_label"]
                    if not out.get("label_source"):
                        out["label_source"] = os.path.basename(c)
                    break
        try:
            stats, top = _stats_and_examples(series)
            out["stats"] = stats
            out["top_values"] = top
            if not out["examples"]:
                import pandas as pd
                nn = pd.Series(series).dropna().head(10)
                out["examples"] = [{"row": int(i), "value": v}
                                   for i, v in nn.items()]
                out["example_cols"] = ["row", "value"]
        except Exception as e:  # noqa: BLE001
            out["degraded"].append(f"stats failed: {e!r}")

    # codebook / paper snippets
    try:
        out["codebook_files"] = [os.path.basename(f) for f in sorted(glob.glob(
            os.path.join(folder, "*")))
            if re.search(r"codebook|survey|questionnaire|instrument|readme",
                         os.path.basename(f), re.I)]
        out["snippets"] = _snippets(key, [var, out.get("uniform_variable", "")])
    except Exception as e:  # noqa: BLE001
        out["degraded"].append(f"snippets failed: {e!r}")
    return jsonify(out)


def _isna(v):
    try:
        import pandas as pd
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return v is None


# --------------------------------------------------------------- assignments
def _draft_path(key, coder):
    return os.path.join(S.CODINGS, f"{key} - assignments - {safe_token(coder)}.json")


def _read_draft(path):
    try:
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return None
    return data if isinstance(data, dict) else None


def _clean_assignment(a):
    if not isinstance(a, dict):
        return None
    row = {k: ("" if a.get(k) is None else str(a.get(k)).strip())
           for k in _DRAFT_FIELDS}
    if not row["original_variable"]:
        return None
    if row["status"] not in ("proposed", "confirmed", "conflict"):
        row["status"] = "proposed"
    return row


def _draft_rows(key):
    """All drafts for ``key`` as (coder, row) pairs, plus the merged view."""
    out, by_coder = [], {}
    for f in glob.glob(os.path.join(S.CODINGS, glob.escape(key) + " - assignments - *.json")):
        base = os.path.basename(f)[: -len(".json")]
        coder = base[len(key) + len(" - assignments - "):]
        data = _read_draft(f)
        if not data:
            continue
        rows = [r for r in (_clean_assignment(a) for a in (data.get("assignments") or []))
                if r]
        by_coder[coder] = rows
        out.extend({"coder": coder, **r} for r in rows)
    return out, by_coder


@bp.get("/api/assignments/<path:key>")
def api_assignments(key):
    if not _known_study(key):
        abort(404)
    merged, by_coder = _draft_rows(key)
    try:
        va = _import_assignments()
        saved = va.load(key, _study_folder(key))
        applied = [{c: _s(v) for c, v in r.items()} for _, r in saved.iterrows()]
    except Exception:  # noqa: BLE001
        applied = []
    return jsonify({"key": key, "drafts": merged, "by_coder": by_coder,
                    "applied": applied, "applied_path": f"{key} - variable_assignments.csv",
                    "fields": list(_DRAFT_FIELDS)})


@bp.post("/api/assignments/<path:key>")
def api_assignments_save(key):
    """Save coder assignments **directly into the canonical dataset file**.

    Accepts ``{"assignments": [...], "coder"?: str, "apply_to_dataset"?: bool}``.
    The coder name is OPTIONAL — a blank name is recorded under the neutral
    label ``coder``.  Every incoming row is upserted into
    ``Datasets/<Key>/<Key> - variable_assignments.csv`` via
    ``variable_assignments.write_rows`` and, unless ``apply_to_dataset`` is
    false, applied to the harmonized sidecar/wide with ``apply_files`` (which
    fills blanks / confirmed overrides / split rows and appends the per-field
    audit trail to ``Results/variable_assignment_log.jsonl``).  Returns the
    resulting changes so the UI can refresh.  A per-coder draft copy is kept in
    ``Codings/<Key> - assignments - <coder>.json`` for audit/back-compat.
    """
    if not _known_study(key):
        return jsonify({"ok": False, "error": "unknown study key"}), 404
    payload = request.get_json(force=True, silent=True) or {}
    coder = str(payload.get("coder") or "").strip() or "coder"
    incoming = payload.get("assignments")
    do_apply = payload.get("apply_to_dataset")
    do_apply = True if do_apply is None else bool(do_apply)
    if not isinstance(incoming, list):
        return jsonify({"ok": False, "error": "assignments must be a list"}), 400
    clean = [r for r in (_clean_assignment(a) for a in incoming) if r]
    if not clean:
        return jsonify({"ok": False, "error": "no assignments to save"}), 400

    try:
        va = _import_assignments()
    except Exception as e:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"variable_assignments unavailable: {e!r}"}), 500

    folder = _study_folder(key)
    stamp = datetime.now().isoformat(timespec="seconds")
    out_rows = []
    for r in clean:
        row = {c: "" for c in va.ASSIGN_COLS}
        for f in _DRAFT_FIELDS:
            if f in row:
                row[f] = r.get(f, "")
        row["coder"] = coder
        row["decided_at"] = stamp
        if not row.get("status"):
            row["status"] = "proposed"
        out_rows.append(row)

    # 1) write the canonical study file (upsert by raw key)
    try:
        n_total = va.write_rows(key, folder, out_rows)
    except Exception as e:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"write_rows failed: {e!r}"}), 500

    # 2) regenerate the harmonized sidecar/wide from the canonical file
    records = []
    if do_apply:
        try:
            records = va.apply_files(key, folder)
        except Exception as e:  # noqa: BLE001
            return jsonify({"ok": False, "error": f"apply_files failed: {e!r}",
                            "rows_written": n_total}), 500

    # 3) keep the per-coder draft copy in sync (audit trail / back-compat).
    # The canonical dataset write above is the contract, so a draft-write problem
    # is reported but never fails the save.
    draft_path = None
    draft_error = None
    try:
        target = _draft_path(key, coder)
        os.makedirs(os.path.dirname(target), exist_ok=True)
        existing = _read_draft(target) or {"key": key, "coder": coder, "assignments": []}
        old = [r for r in (_clean_assignment(a) for a in (existing.get("assignments") or [])) if r]
        incoming_keys = {(r["artifact"], r["original_variable"], r["item_scope"]) for r in clean}
        keep = [r for r in old
                if (r["artifact"], r["original_variable"], r["item_scope"]) not in incoming_keys]
        merged = keep + clean
        record = {"key": key, "coder": coder, "saved": stamp, "assignments": merged}
        fd, tmp = tempfile.mkstemp(suffix=".json", dir=S.TEMP)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as fh:
                json.dump(record, fh, ensure_ascii=False, indent=2)
                fh.write("\n")
            atomic_replace(tmp, target)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        merged_total = len(merged)
        draft_path = os.path.relpath(target, S.ROOT).replace("\\", "/")
    except Exception as e:  # noqa: BLE001
        merged_total = len(clean)
        draft_error = repr(e)

    actions = [{"action": r.get("action"), "original_variable": r.get("original_variable"),
                "field": r.get("field"), "old": r.get("old"), "new": r.get("new"),
                "status": r.get("status"), "note": r.get("note")}
               for r in records]
    return jsonify({"ok": True, "coder": coder, "count": len(clean), "total": merged_total,
                    "rows_written": n_total, "applied_to_dataset": do_apply,
                    "changes": len(actions), "records": actions,
                    "path": f"{key} - variable_assignments.csv",
                    "draft_path": draft_path, "draft_error": draft_error})


@bp.post("/api/assignments/<path:key>/apply")
def api_assignments_apply(key):
    """Apply already-saved drafts to ``<Key> - variable_assignments.csv`` and
    (unless ``apply_to_dataset`` is false) regenerate the sidecar/wide.

    Kept for backwards compatibility; the assignment workspace now saves through
    ``POST /api/assignments/<key>`` in one sweep."""
    if not _known_study(key):
        return jsonify({"ok": False, "error": "unknown study key"}), 404
    payload = request.get_json(force=True, silent=True) or {}
    coder = str(payload.get("coder") or "").strip()
    do_apply = payload.get("apply_to_dataset")
    do_apply = True if do_apply is None else bool(do_apply)
    folder = _study_folder(key)
    merged, by_coder = _draft_rows(key)
    if coder:
        rows = by_coder.get(coder) or []
    else:
        rows = [r for r in merged]
    if not rows:
        return jsonify({"ok": False, "error": "no assignment drafts to apply"}), 400

    try:
        va = _import_assignments()
    except Exception as e:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"variable_assignments unavailable: {e!r}"}), 500

    stamp = datetime.now().isoformat(timespec="seconds")
    out_rows = []
    for r in rows:
        row = {c: "" for c in va.ASSIGN_COLS}
        for f in _DRAFT_FIELDS:
            if f in row:
                row[f] = r.get(f, "")
        row["coder"] = coder or r.get("coder", "")
        row["decided_at"] = stamp
        if not row.get("status"):
            row["status"] = "proposed"
        out_rows.append(row)
    try:
        n_total = va.write_rows(key, folder, out_rows)
    except Exception as e:  # noqa: BLE001
        return jsonify({"ok": False, "error": f"write_rows failed: {e!r}"}), 500

    records = []
    if do_apply:
        try:
            records = va.apply_files(key, folder)
        except Exception as e:  # noqa: BLE001
            return jsonify({"ok": False, "error": f"apply_files failed: {e!r}",
                            "rows_written": n_total}), 500
    actions = [{"action": r.get("action"), "original_variable": r.get("original_variable"),
                "field": r.get("field"), "old": r.get("old"), "new": r.get("new"),
                "status": r.get("status"), "note": r.get("note")}
               for r in records]
    return jsonify({"ok": True, "rows_written": n_total, "applied_to_dataset": do_apply,
                    "changes": len(actions), "records": actions,
                    "path": f"{key} - variable_assignments.csv"})
